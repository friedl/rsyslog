
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()

    # Slide 1: Title
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Adiscon GmbH – Excellence in Logging & System Management"
    subtitle.text = "Innovative Lösungen aus Großrinderfeld für eine weltweit vernetzte IT-Infrastruktur.\n\nPräsentator: [Dein Name/Position]"

    # Slide 2: Firmenprofil
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    slide.shapes.title.text = "Firmenprofil Adiscon GmbH"
    tf = slide.placeholders[1].text_frame
    tf.text = "Standort: Großrinderfeld, Deutschland"
    p = tf.add_paragraph()
    p.text = "Historie: Gegründet 1988 durch Rainer Gerhards"
    p = tf.add_paragraph()
    p.text = "Philosophie: Customer-Driven Innovation"
    p = tf.add_paragraph()
    p.text = "Globaler Einfluss: Macher hinter rsyslog (Industriestandard)"

    # Slide 3: rsyslog
    slide = prs.slides.add_slide(bullet_slide_layout)
    slide.shapes.title.text = "Kernprodukt: rsyslog (Open Source)"
    tf = slide.placeholders[1].text_frame
    tf.text = "Extrem schneller, modularer Log-Processing-Dienst"
    p = tf.add_paragraph()
    p.text = "Leistung: > 1 Mio. Nachrichten/Sekunde"
    p = tf.add_paragraph()
    p.text = "Protokolle: Unterstützt RELP für verlustfreie Übertragung"
    p = tf.add_paragraph()
    p.text = "Verbreitung: Standard in fast allen Linux-Distributionen"

    # Slide 4: Windows-Lösungen
    slide = prs.slides.add_slide(bullet_slide_layout)
    slide.shapes.title.text = "Windows-Lösungen: WinSyslog & EventReporter"
    tf = slide.placeholders[1].text_frame
    tf.text = "WinSyslog: Erster Syslog-Server für Windows (seit 1996)"
    p = tf.add_paragraph()
    p.text = "EventReporter: Schlankes Monitoring von Windows Event Logs"
    p = tf.add_paragraph()
    p.text = "Fokus: Integration von Windows in heterogene Umgebungen"

    # Slide 5: Enterprise-Lösungen
    slide = prs.slides.add_slide(bullet_slide_layout)
    slide.shapes.title.text = "Enterprise: MonitorWare Agent & Rsyslog Windows Agent"
    tf = slide.placeholders[1].text_frame
    tf.text = "MonitorWare Agent: All-in-One Lösung (Input/Output/Processing)"
    p = tf.add_paragraph()
    p.text = "Rsyslog Windows Agent: Native RELP-Unterstützung für Windows"
    p = tf.add_paragraph()
    p.text = "Einsatz: Hochperformante Serverfarmen & kritische Infrastrukturen"

    # Slide 6: Datenfluss-Diagramm
    slide = prs.slides.add_slide(prs.slide_layouts[5]) # Title Only layout
    slide.shapes.title.text = "Datenfluss & Produktvernetzung"

    # Adding shapes for diagram
    def add_box(slide, text, left, top, width=1.5, height=0.8, color=RGBColor(0, 102, 204)):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.text = text
        # Set font size
        tf = shape.text_frame
        for p in tf.paragraphs:
            p.font.size = Pt(10)
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        return shape

    # Source Boxes
    box_win_agent = add_box(slide, "Rsyslog Windows Agent\n(Windows Server)", 0.5, 1.5, 2.0, 1.0)
    box_evnt_rep = add_box(slide, "EventReporter\n(Windows Clients)", 0.5, 3.0, 2.0, 1.0)
    box_linux = add_box(slide, "Linux / Unix\n(Native Logging)", 0.5, 4.5, 2.0, 1.0)

    # Central Hub
    box_rsyslog = add_box(slide, "rsyslog Server\n(Zentraler Hub / Linux)", 4.0, 2.5, 2.0, 1.5, RGBColor(204, 0, 0))
    box_winsyslog = add_box(slide, "WinSyslog\n(Central Hub / Win)", 4.0, 4.5, 2.0, 1.0, RGBColor(0, 153, 0))

    # Analysis/Output
    box_mw_agent = add_box(slide, "MonitorWare Agent\n(Analyse / DB / Alarm)", 7.5, 2.75, 2.0, 1.0, RGBColor(255, 153, 0))

    # Connectors (simplistic arrows)
    def add_arrow(slide, start_x, start_y, end_x, end_y):
        connector = slide.shapes.add_connector(1, Inches(start_x), Inches(start_y), Inches(end_x), Inches(end_y))
        return connector

    # Arrows from sources to rsyslog/winsyslog
    add_arrow(slide, 2.5, 2.0, 4.0, 3.0) # Win Agent -> rsyslog
    add_arrow(slide, 2.5, 3.5, 4.0, 3.25) # EventRep -> rsyslog
    add_arrow(slide, 2.5, 5.0, 4.0, 5.0) # Linux -> WinSyslog

    # Connection between hubs
    add_arrow(slide, 5.0, 4.0, 5.0, 4.5) # rsyslog -> WinSyslog

    # Arrow to MonitorWare
    add_arrow(slide, 6.0, 3.25, 7.5, 3.25) # rsyslog -> MW Agent

    # Add text labels for protocols
    def add_label(slide, text, left, top):
        txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(1.5), Inches(0.5))
        tf = txBox.text_frame
        tf.text = text
        for p in tf.paragraphs:
            p.font.size = Pt(9)

    add_label(slide, "RELP / TLS", 2.8, 2.0)
    add_label(slide, "Syslog / UDP", 2.8, 3.5)
    add_label(slide, "RELP", 6.5, 3.0)

    # Slide 7: Fazit
    slide = prs.slides.add_slide(bullet_slide_layout)
    slide.shapes.title.text = "Fazit & Vorteile"
    tf = slide.placeholders[1].text_frame
    tf.text = "Skalierbarkeit: Von KMU bis Enterprise"
    p = tf.add_paragraph()
    p.text = "Zuverlässigkeit: Wegbereiter für RELP"
    p = tf.add_paragraph()
    p.text = "Expertise: Direkter Support durch Entwickler"
    p = tf.add_paragraph()
    p.text = "Qualität: Made in Germany (Großrinderfeld)"

    prs.save('Adiscon_Unternehmenspraesentation.pptx')

if __name__ == "__main__":
    create_presentation()
